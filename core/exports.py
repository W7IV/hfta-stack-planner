"""Exports: HFTA replay pack (zip), Excel workbook with traceable formulas, CSV, self-contained HTML report."""
from __future__ import annotations
import csv, datetime, hashlib, html, io, json, zipfile
from pathlib import Path
import numpy as np
from . import antennas as A, project as P, terrain as TR

def sha256_file(f) -> str:
    return hashlib.sha256(Path(f).read_bytes()).hexdigest()

def names(proj):
    return ({s["id"]: s["name"] for s in proj["sites"]}, {t["id"]: t["name"] for t in proj["towers"]},
            {t["id"]: t["name"] for t in proj["targets"]}, {t["id"]: t for t in proj["targets"]})

def run_rows(proj, res):
    """One row per (HFTA run, elevation file, band) actually used."""
    S, _, _, TG = names(proj); rows = []; notes = res.get("run_notes") or {}
    for rid, spec in res["runs"].items():
        for key, fom in res["fom_log"].get(rid, {}).items():
            prn, band = key.split("|")
            hs = spec["heights_ft"] + [None] * (4 - len(spec["heights_ft"]))
            rows.append({"key": f"{rid}|{prn}|{band}", "run_id": rid, "site": S[spec["site"]], "terrain_file": spec["file"],
                         "elevation_file": prn, "band_m": band, "freq_mhz": spec["freq"], "antenna_type": spec["hfta_type"],
                         "n_antennas": len(spec["heights_ft"]), "h1_ft": hs[0], "h2_ft": hs[1], "h3_ft": hs[2], "h4_ft": hs[3],
                         "phasing": "all in phase", "dielectric": spec["epsr"], "conductivity_S_m": spec["sigma"],
                         "diffraction": "on" if spec["diffraction"] else "off", "hfta_fom_db": fom,
                         "terrain_sha256": spec.get("terrain_sha256", ""),
                         "core_note": ("filled by interpolation at " + ", ".join(f"{a:g}" for a in notes[rid]["filled_deg"]) + " deg")
                                      if (notes.get(rid) or {}).get("filled_deg") else (notes.get(rid) or {}).get("failed", "")})
    return sorted(rows, key=lambda r: (r["site"], r["terrain_file"], r["freq_mhz"], r["key"]))

def cell_rows(proj, res, layers_on):
    S, TW, TN, _ = names(proj); out = []
    from .evaluate import total
    for c in res["cells"]:
        best = c["options"][c["best"]] if c.get("best") is not None else None
        out.append({"site": S[c["site"]], "tower": TW[c["tower"]], "target": TN[c["target"]], "band_m": c["band"], "freq_mhz": c["freq"],
                    "antennas": " / ".join(f'{a["h"]:g} ft' for a in c["antennas"]),
                    "hfta_type": c.get("hfta_type", ""),
                    "switch_setting": " / ".join(f"{h:g}" for h in best["heights"]) if best else "",
                    "hfta_fom_db": c.get("hfta_fom"),
                    "spacing_warning": "; ".join(f'{w["level"].upper()}: {w["text"]}' for w in c.get("spacing", [])),
                    "issues": "; ".join(c["issues"]), "notes": "; ".join(c.get("notes", []))})
    return out

RECIPE = """# Reproducing these results in the original HFTA

Calculations by OpenHFTA (RioDXGroup, https://github.com/RioDXGroup/openhfta), an open-source implementation of the
HFTA method, used unchanged; the engine version is in `manifest.json`. HFTA by R. Dean Straw, N6BV, published with
The ARRL Antenna Book. ARRL elevation-angle files are copyright ARRL. Not affiliated with or endorsed by ARRL or the
OpenHFTA authors.

Every Figure of Merit (FOM) in this study is either one OpenHFTA run or simple arithmetic on runs. OpenHFTA agrees
with original-HFTA output within 0.1 dB on the reference set in the planner's tests; other cases are not proven
identical. Spacing warnings and the guy check are notes only; they never change a number.

## Files in this pack
- `terrain/<site>/` : the terrain profiles (.PRO) exactly as used (SHA-256 in `manifest.json` and `runs.csv`).
- `prn/` : the elevation-angle (.PRN) files, ONLY in a private pack. A shareable pack leaves them out: copyright
  ARRL (or the user's own data). `manifest.json` lists each file by name and SHA-256, so a recipient can check that
  their own copy is the same file before comparing FOMs.
- `runs.csv` : every run with all its inputs and the FOM it gives. `core_note`: a run where the calculation core
  returned a few invalid angles that were filled by interpolation (not an exact HFTA result), or a failed run.
- `results.csv` : every reported result, with the switch setting chosen.
- `combination.csv` : how each result is built from runs (headings, weights, switch settings).
- `expected_out/<run_id>.PRN` : the elevation pattern (gain at 1-35 degrees) each run should give, in the same
  layout as HFTA's OUT.PRN, for a direct comparison.

## One HFTA run (one row of runs.csv)
1. Start HFTA. Set units to feet.
2. Frequency: `freq_mhz`.
3. For each antenna 1 to `n_antennas`: tick it, enter its height (`h1_ft` ...), choose the antenna type
   `antenna_type`, and load the same terrain file `terrain_file` for every antenna. All antennas in phase.
4. Elevation-angle file: load `elevation_file`. HFTA takes the column for the band of the frequency.
5. Ground: dielectric constant `dielectric`, conductivity `conductivity_S_m`. Diffraction: `diffraction`.
6. Compute. The FOM shown by HFTA is `hfta_fom_db`. The pattern can be compared with `expected_out/<run_id>.PRN`.

## From runs to a result
- Direction: HFTA treats the antennas as facing the azimuth of the terrain file. A tower is evaluated only on the
  targets chosen for it, so the terrain files set the direction.
- Headings: a target is evaluated at one or more headings (e.g. -8, 0, +8 degrees around the target bearing). Each
  heading uses the nearest terrain file. The heading value = weighted mean of the run FOMs in dB:
  sum(weight x FOM) / sum(weight). A switch setting with a failed run has no value.
- Switch setting: when the tower has a stack switch, each setting (whole stack, pairs, single antennas) is a
  separate set of runs. The reported value is the highest setting. All settings are listed in combination.csv.
- Bands and targets are always reported separately. There are no hidden totals.
"""

def replay_pack(proj, pslug, res, layers_on, include_prn: bool = False) -> bytes:
    """include_prn=False (default): shareable pack without .PRN files (listed by name and SHA-256 in manifest.json).
    include_prn=True: private pack with the user's own .PRN files, for the owner's own records."""
    S, TW, TN, TG = names(proj); buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        kind = "PRIVATE (contains your .PRN files: do not share)" if include_prn else "SHAREABLE (no .PRN files)"
        z.writestr("README_HFTA.md", RECIPE + f"\nProject: {proj['name']}. Exported {datetime.date.today()}. Pack: {kind}.\n")
        used_files = {(s["site"], s["file"]) for s in res["runs"].values()}
        sites_ok = {s["id"] for s in proj["sites"]}
        terrain = []
        for sid, f in sorted(used_files):
            src = P.site_dir(pslug, sid) / f
            if sid in sites_ok and src.exists():
                z.write(src, f"terrain/{S[sid]}/{f}"); terrain.append({"site": S[sid], "file": f, "sha256": sha256_file(src)})
        prn_list = []
        for prn in sorted({t["prn"] for t in proj["targets"] if t.get("prn")}):
            src = P.prn_file(pslug, prn)
            if not src.exists(): continue
            now = sha256_file(src); then = (res.get("prn_sha256") or {}).get(prn)
            prn_list.append({"file": prn, "sha256": now, "included": include_prn,
                             "changed_since_run": bool(then and then != now)})
            if include_prn: z.write(src, f"prn/{prn}")
        z.writestr("manifest.json", json.dumps({"pack": "private" if include_prn else "shareable", "project": proj["name"],
                   "results_stamp": res.get("stamp"), "engine": res.get("engine") or {"name": "OpenHFTA", "commit": "unknown"},
                   "prn_files": prn_list, "terrain_files": terrain,
                   "note": "Elevation-angle (.PRN) files are not included in a shareable pack. Use your own copy; compare its SHA-256."}, indent=1))
        rr = run_rows(proj, res); z.writestr("runs.csv", _csv(rr))
        z.writestr("results.csv", _csv(cell_rows(proj, res, layers_on)))
        comb = []
        for c in res["cells"]:
            for oi, o in enumerate(c["options"]):
                for a in o["az"]:
                    comb.append({"site": S[c["site"]], "tower": TW[c["tower"]], "target": TN[c["target"]], "band_m": c["band"],
                                 "switch_setting": " / ".join(f"{h:g}" for h in o["heights"]), "setting_score_db": o.get("score"),
                                 "chosen": oi == c.get("best"), "heading_deg": a["az"], "terrain_heading_deg": a["file_az"],
                                 "weight": a["weight"], "run_id": a["run"], "run_fom_db": a.get("fom")})
        z.writestr("combination.csv", _csv(comb))
        for rid, pat in res["patterns"].items():
            if pat is None: continue
            spec = res["runs"][rid]
            lines = [f'" {spec["freq"]:4.1f} MHz    Ants: {"/".join(f"{h:g}" for h in spec["heights_ft"])} feet"',
                     f'"HFTA: {spec["file"]}"', f'"Antenna: {spec["hfta_type"]}."', '"  Elev."," Gain "']
            lines += [f"{e}, {pat[e * 4 - 1]:.1f}" for e in range(1, 36)]
            z.writestr(f"expected_out/{rid}.PRN", "\n".join(lines) + "\n")
        z.writestr("project.json", json.dumps(proj, indent=1))
    return buf.getvalue()

_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")

def _cell_text(v):
    """CSV for spreadsheets: text that a spreadsheet would read as a formula gets a leading apostrophe."""
    return "'" + v if isinstance(v, str) and v.startswith(_FORMULA_START) else v

def _csv(rows):
    if not rows: return ""
    s = io.StringIO(); w = csv.DictWriter(s, fieldnames=list(rows[0])); w.writeheader()
    w.writerows({k: _cell_text(v) for k, v in r.items()} for r in rows); return s.getvalue()

def results_csv(proj, res, layers_on): return _csv(cell_rows(proj, res, layers_on))

# ---------------------------------------------------------------- Excel
def workbook(proj, pslug, res, layers_on, guy_rows=None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter as L
    F = lambda **k: Font(name="Arial", **k)
    HEAD, BLUE, BOLD, ITAL = F(bold=True, color="FFFFFF"), F(color="0000FF"), F(bold=True), F(italic=True)
    HF = PatternFill("solid", fgColor="23395B"); th = Side(style="thin", color="C9D3DF"); BOX = Border(top=th, bottom=th, left=th, right=th)
    def hdr(ws, r, vals):
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v); c.font = HEAD; c.fill = HF; c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r].height = 30
    def put(ws, r, c, v, font=None, fmt=None, formula=False):
        """Text is always stored as text, so a name such as '=1+1' is never run as a formula.
        Only the workbook's own formulas (formula=True) are formulas."""
        x = ws.cell(r, c); x.value = v
        if isinstance(v, str) and not formula: x.data_type = "s"
        x.font = font or F(); x.border = BOX
        if fmt: x.number_format = fmt
    def text(ws, ref, v):
        ws[ref] = v
        if isinstance(v, str): ws[ref].data_type = "s"
        return ws[ref]
    S, TW, TN, TG = names(proj); wb = Workbook()
    # Runs
    wr = wb.active; wr.title = "Runs"
    wr["A1"] = "HFTA runs: one row = one run in the original HFTA (inputs and the FOM it gives). See README in the replay pack."; wr["A1"].font = BOLD
    rr = run_rows(proj, res); cols = list(rr[0]) if rr else []
    hdr(wr, 3, cols)
    for i, row in enumerate(rr, 4):
        for j, k in enumerate(cols, 1): put(wr, i, j, row[k], BLUE if k == "hfta_fom_db" else None, "0.00" if k == "hfta_fom_db" else None)
    nr = 3 + len(rr)
    # Combination (per heading use)
    wc = wb.create_sheet("Combination")
    wc["A1"] = "How each result is built: one row per heading of each switch setting. FOM is looked up from Runs."; wc["A1"].font = BOLD
    hdr(wc, 3, ["Setting key", "Cell key", "Heading (deg)", "Terrain file heading", "Weight", "Run key", "Run FOM (dB)", "Weight x FOM", "Weight used"])
    r = 4; opt_rows = []
    for ci, c in enumerate(res["cells"]):
        if c["issues"]: continue
        tg = TG[c["target"]]
        for oi, o in enumerate(c["options"]):
            okey = f"C{ci}-S{oi}"; opt_rows.append((okey, f"C{ci}", c, o, oi))
            for a in o["az"]:
                put(wc, r, 1, okey); put(wc, r, 2, f"C{ci}"); put(wc, r, 3, a["az"]); put(wc, r, 4, a["file_az"]); put(wc, r, 5, a["weight"], BLUE)
                put(wc, r, 6, f'{a["run"]}|{tg["prn"]}|{c["band"]}')
                put(wc, r, 7, f'=IFERROR(IF(INDEX(Runs!$R$4:$R${nr},MATCH(F{r},Runs!$A$4:$A${nr},0))="","",INDEX(Runs!$R$4:$R${nr},MATCH(F{r},Runs!$A$4:$A${nr},0))),"")', None, "0.00", True)
                put(wc, r, 8, f'=IF(G{r}="","",E{r}*G{r})', None, "0.00", True); put(wc, r, 9, f'=IF(G{r}="",0,E{r})', formula=True); r += 1
    nc = r - 1
    # Settings
    ws_ = wb.create_sheet("Settings")
    ws_["A1"] = "Switch settings: weighted mean over headings (dB). The chosen setting of each result is the highest."; ws_["A1"].font = BOLD
    hdr(ws_, 3, ["Setting key", "Cell key", "Site", "Tower", "Target", "Band (m)", "Antennas in use (ft)", "Score (dB)"])
    for i, (okey, ckey, c, o, oi) in enumerate(opt_rows, 4):
        put(ws_, i, 1, okey); put(ws_, i, 2, ckey); put(ws_, i, 3, S[c["site"]]); put(ws_, i, 4, TW[c["tower"]]); put(ws_, i, 5, TN[c["target"]])
        put(ws_, i, 6, c["band"]); put(ws_, i, 7, " / ".join(f"{h:g}" for h in o["heights"]))
        # No value when any heading of the setting has no FOM (failed run): same rule as the planner.
        put(ws_, i, 8, f'=IF(COUNTIFS(Combination!$A$4:$A${nc},A{i},Combination!$G$4:$G${nc},"")>0,"",IFERROR(SUMIFS(Combination!$H$4:$H${nc},Combination!$A$4:$A${nc},A{i})/SUMIFS(Combination!$I$4:$I${nc},Combination!$A$4:$A${nc},A{i}),""))', None, "0.00", True)
    ns = 3 + len(opt_rows)
    # Results
    wres = wb.create_sheet("Results", 0)
    text(wres, "A1", f"{proj['name']}: results (FOM, dB), calculated with OpenHFTA (HFTA method).").font = F(bold=True, size=12)
    wres["A2"] = "Spacing warning: RED = too close, AMBER = close. HFTA has no mutual coupling; a warning never changes the FOM."; wres["A2"].font = ITAL
    hdr(wres, 4, ["Cell key", "Site", "Tower", "Target", "Band (m)", "Freq (MHz)", "Antenna heights (ft)", "HFTA type", "HFTA FOM (dB)",
                  "Spacing warning", "Note"])
    for i, (ci, c) in enumerate([(ci, c) for ci, c in enumerate(res["cells"])], 5):
        put(wres, i, 1, f"C{ci}"); put(wres, i, 2, S[c["site"]]); put(wres, i, 3, TW[c["tower"]]); put(wres, i, 4, TN[c["target"]])
        put(wres, i, 5, c["band"]); put(wres, i, 6, c["freq"]); put(wres, i, 7, " / ".join(f'{a["h"]:g}' for a in c["antennas"]))
        put(wres, i, 8, c.get("hfta_type", ""))
        sp = c.get("spacing", [])
        put(wres, i, 10, "; ".join(f'{w["level"].upper()}: {w["text"]}' for w in sp),
            F(color="B33A3A", bold=True) if any(w["level"] == "red" for w in sp) else F(color="B26B00") if sp else None)
        if c["issues"]:
            put(wres, i, 11, "; ".join(c["issues"])); continue
        put(wres, i, 9, f'=IF(COUNTIFS(Settings!$B$4:$B${ns},A{i},Settings!$H$4:$H${ns},">-1E+300")=0,"",_xlfn.MAXIFS(Settings!$H$4:$H${ns},Settings!$B$4:$B${ns},A{i}))', BOLD, "0.00", True)
    for ws in (wres, wr, wc, ws_):
        for j in range(1, 20): ws.column_dimensions[L(j)].width = 14
    wres.column_dimensions["G"].width = 24; wres.column_dimensions["J"].width = 70; wres.column_dimensions["K"].width = 50; wres.freeze_panes = "B5"
    # Sites
    wsi = wb.create_sheet("Sites")
    hdr(wsi, 1, ["Site", "Latitude", "Longitude", "Terrain source", "Profiles", "Notes"])
    for i, s in enumerate(proj["sites"], 2):
        put(wsi, i, 1, s["name"]); put(wsi, i, 2, s.get("lat")); put(wsi, i, 3, s.get("lon")); put(wsi, i, 4, s.get("terrain", {}).get("source"))
        put(wsi, i, 5, len(TR.profile_index(P.site_dir(pslug, s["id"])))); put(wsi, i, 6, s.get("notes", ""))
    # Inputs
    wi = wb.create_sheet("Inputs")
    st = proj["settings"]; r = 1
    for k, v in (("Project", proj["name"]), ("Ground", f'{st["soil"]}: dielectric {st["epsr"]}, conductivity {st["sigma"]} S/m'),
                 ("Diffraction", "on" if st["diffraction"] else "off"),
                 ("Engine", "OpenHFTA commit {commit}, library SHA-256 {library_sha256}, wrapper {wrapper}".format(
                     **{"commit": "unknown", "library_sha256": "unknown", "wrapper": "?", **(res.get("engine") or {})})),
                 ("Results computed", res.get("stamp", ""))):
        wi.cell(r, 1, k).font = BOLD; x = wi.cell(r, 2); x.value = v; x.data_type = "s"; r += 1
    r += 1; hdr(wi, r, ["Target", "Elevation file (.PRN)", "Heading (deg)", "Offsets (deg)", "Weights", "Bands (m)"]); r += 1
    for t in proj["targets"]:
        for j, v in enumerate([t["name"], t.get("prn") or "(none)", t["heading"], ", ".join(map(str, t["offsets"])), ", ".join(map(str, t["weights"])), ", ".join(t["bands"])], 1):
            put(wi, r, j, v)
        r += 1
    sp = st.get("spacing") or {}
    r += 1; wi.cell(r, 1, "Spacing warnings").font = BOLD
    wi.cell(r, 2, f'stack: red below {sp.get("red_wl")} wavelength, amber below {sp.get("amber_wl")}; gap between antennas on different bands: red below {sp.get("gap_red_ft")} ft, amber below {sp.get("gap_amber_ft")} ft'); r += 1
    r += 1; hdr(wi, r, ["Tower", "Serves targets", "Switching", "Antenna (band and elements, as entered in HFTA)", "Height (ft)", "Turning radius (ft)", "Droop (ft)"]); r += 1
    for t in proj["towers"]:
        for a in sorted(t["antennas"], key=lambda a: -a["h"]):
            serves = ", ".join(TN[x] for x in (t["targets"] if t.get("targets") is not None else TN) if x in TN) or "none"
            for j, v in enumerate([t["name"], serves, t.get("switching"), A.describe(a), a["h"], a.get("turning_radius"), a.get("droop")], 1): put(wi, r, j, v)
            r += 1
    for j in range(1, 7): wi.column_dimensions[L(j)].width = 22
    if guy_rows:
        wg = wb.create_sheet("Guy check")
        wg["A1"] = "Guy clearance (not part of HFTA). Negative = the antenna would touch a guy."; wg["A1"].font = BOLD
        hdr(wg, 3, ["Tower", "Antenna", "Height (ft)", "Turning radius (ft)", "Clearance (ft)", "Limited by"])
        r = 4
        for tname, rows in guy_rows.items():
            for g in rows:
                for j, v in enumerate([tname, g["name"], g["h"], g["turning_radius"], g["clearance"], g["limit"]], 1): put(wg, r, j, v)
                r += 1
        for j, w in zip(range(1, 7), (20, 22, 11, 14, 13, 60)): wg.column_dimensions[L(j)].width = w
    from . import spacing as SP
    from .evaluate import band_freq
    wsp = wb.create_sheet("Spacing check")
    wsp["A1"] = "Spacing warnings (not part of HFTA; they never change a FOM). RED = too close, AMBER = close. Limits: see Inputs."; wsp["A1"].font = BOLD
    hdr(wsp, 3, ["Tower", "Level", "Kind", "Band (m)", "Heights (ft)", "Spacing (ft)", "Spacing (wavelength)", "Warning"])
    r = 4
    for t in proj["towers"]:
        ws_l = SP.check(t["antennas"], lambda bd: band_freq(proj, bd), st["spacing"])
        if not ws_l:
            for j, v in enumerate([t["name"], "OK", "", "", "", "", "", "No spacing warnings"], 1): put(wsp, r, j, v)
            r += 1
        for w in ws_l:
            for j, v in enumerate([t["name"], w["level"].upper(), "stack" if w["kind"] == "stack" else "gap between bands", w["band"] or "",
                                   " / ".join(f"{h:g}" for h in w["heights"]), w["ft"], w["wl"], w["text"]], 1):
                put(wsp, r, j, v, F(color="B33A3A", bold=True) if (j == 2 and w["level"] == "red") else F(color="B26B00") if j == 2 else None)
            r += 1
    for j, wd in zip(range(1, 9), (22, 9, 18, 9, 14, 11, 12, 100)): wsp.column_dimensions[L(j)].width = wd
    b = io.BytesIO(); wb.save(b); return b.getvalue()
