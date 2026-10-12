"""Tests of the planner's own guarantees (not of RF physics). Uses the real OpenHFTA core where a calculation is
needed, the synthetic test .PRN file, and a temporary projects folder. Run: python tests/test_wrapper.py
Each test names the review finding it covers (F01..F22, L03)."""
import io, json, math, os, sys, tempfile, time, zipfile
from pathlib import Path

TMP = Path(tempfile.mkdtemp(prefix="hfta-test-"))
os.environ["PROJECTS_DIR"] = str(TMP / "projects")
os.environ.pop("APP_PASSWORD", None); os.environ.pop("RAILWAY_ENVIRONMENT", None); os.environ.pop("HOSTED", None)
os.environ.setdefault("HFTA_WORKERS", "2")
REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

import numpy as np
import app as APP
import engine
from core import project as P, terrain as TR, evaluate as EV, compromise as CM, guys as G, exports as EX

C = APP.app.test_client()
DATA = REPO / "tests" / "data"; PRN = REPO / "samples" / "TEST-ONLY-synthetic.PRN"
PASSED = []

def test(fn):
    fn(); PASSED.append(fn.__name__); print("ok ", fn.__name__); return fn

def new_project(name="T"):
    slug = C.post("/api/v2/projects", json={"name": name}).get_json()["slug"]
    proj = C.get(f"/api/v2/p/{slug}").get_json()["project"]
    return slug, proj

def put(slug, proj, **kw):
    return C.put(f"/api/v2/p/{slug}", json=proj, **kw)

def setup_study(name="Study", bands=("20",)):
    """One site with the four KK6P reference profiles, one target (synthetic .PRN), one tower with two antennas."""
    slug, proj = new_project(name)
    files = {"files": [(io.BytesIO((DATA / f).read_bytes()), f) for f in ("AZI-25.00.PRO", "AZI-30.00.PRO", "AZI-35.00.PRO", "AZI-40.00.PRO")]}
    proj["sites"] = [{"id": "s1", "name": "Site 1"}]; put(slug, proj)
    r = C.post(f"/api/v2/p/{slug}/sites/s1/import", data=files, content_type="multipart/form-data"); assert r.status_code == 200, r.data
    C.post(f"/api/v2/p/{slug}/prn", data={"files": [(io.BytesIO(PRN.read_bytes()), "SYN.PRN")]}, content_type="multipart/form-data")
    proj = C.get(f"/api/v2/p/{slug}").get_json()["project"]
    proj["targets"] = [{"id": "t1", "name": "Europe", "prn": "SYN.PRN", "heading": 30, "offsets": [-5, 0, 5], "weights": [1, 2, 1], "bands": list(bands)}]
    proj["towers"] = [{"id": "w1", "name": "Tower", "switching": "stackmatch", "antennas": [{"h": 60, "el": {b: 4 for b in bands}}, {"h": 120, "el": {b: 4 for b in bands}}]}]
    r = put(slug, proj); assert r.status_code == 200, r.data
    return slug, r.get_json()["project"]

def evaluate(slug):
    jid = C.post(f"/api/v2/p/{slug}/evaluate").get_json()["job"]
    for _ in range(600):
        j = C.get(f"/api/job/{jid}").get_json()
        if j["state"] != "running": break
        time.sleep(0.2)
    assert j["state"] == "done", j
    return j["result"], json.loads((P.path(slug) / "results.json").read_text())

# ------------------------------------------------------------------ F01 path containment
@test
def f01_ids_and_paths_are_contained():
    slug, proj = new_project("Paths")
    outside = TMP / "outside"; outside.mkdir(); (outside / "sentinel.txt").write_text("secret")
    for bad in ("../outside", str(outside), "..", "a/b", "C:\\x", "s 1"):
        p2 = json.loads(json.dumps(proj)); p2["sites"] = [{"id": bad, "name": "x"}]
        assert put(slug, p2).status_code == 400, bad
    for bad in ("../../etc/passwd.prn", "..prn", "/abs.prn"):
        p2 = json.loads(json.dumps(proj)); p2["targets"] = [{"id": "t1", "name": "x", "prn": bad, "heading": 0, "bands": ["20"]}]
        assert put(slug, p2).status_code == 400, bad
    assert C.get(f"/api/v2/p/{slug}/prn/..%2F..%2Fproject.json").status_code in (400, 404)
    assert C.get("/api/v2/p/..%2Foutside").status_code in (400, 404)
    # an imported zip with a path that climbs out of the project, and an absolute site id
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr("project.json", json.dumps({"name": "Evil", "sites": [{"id": "..", "name": "x"}], "targets": [], "towers": []}))
        z.writestr("sites/../escaped.PRO", "meters\n0 0\n30 1\n")
    r = C.post("/api/v2/projects/import", data={"file": (io.BytesIO(b.getvalue()), "e.zip")}, content_type="multipart/form-data")
    assert r.status_code == 400
    assert not list(Path(os.environ["PROJECTS_DIR"]).rglob("escaped.PRO"))
    assert (outside / "sentinel.txt").read_text() == "secret"

# ------------------------------------------------------------------ F02 types checked on save and import
@test
def f02_markup_in_numbers_is_refused():
    slug, proj = new_project("Xss")
    proj["settings"]["epsr"] = '13"><img src=x onerror=alert(1)>'
    assert put(slug, proj).status_code == 400
    proj["settings"]["epsr"] = 13; proj["sites"] = [{"id": "s1", "name": "<b>ok as text</b>", "lat": "<script>"}]
    assert put(slug, proj).status_code == 400

# ------------------------------------------------------------------ F13 terrain rows and weights
@test
def f13_profile_rows_parsed_together():
    d, h = engine.parse_profile_text("feet\n0 0\n100 bad\n200 20\n")
    assert d == [0.0, 200.0] and h == [0.0, 20.0], (d, h)
    for bad in ("feet\n0 0\n200 5\n100 6\n", "feet\n0 0\n100 nan\n", "feet\n0 0\n100 inf\n", "feet\n0 0\n0 5\n"):
        try: engine.parse_profile_text(bad)
        except ValueError: continue
        raise AssertionError(f"accepted: {bad!r}")
    slug, proj = new_project("Weights")
    proj["targets"] = [{"id": "t1", "name": "x", "prn": "", "heading": 0, "offsets": [0], "weights": [0], "bands": ["20"]}]
    assert put(slug, proj).status_code == 400
    proj["targets"][0].update(offsets=[0, 5], weights=[1])
    assert put(slug, proj).status_code == 400

# ------------------------------------------------------------------ F05 invalid core output
@test
def f05_invalid_core_output_policy():
    base = np.linspace(-10, 10, 140)
    few = base.copy(); few[77:82] = np.nan                       # the known core fault: 5 angles (1.25 deg)
    pat, note = EV.check_pattern(few)
    assert pat is not None and np.isfinite(pat).all() and len(note["filled_deg"]) == 5
    many = base.copy(); many[2:140] = np.nan
    assert EV.check_pattern(many) == (None, EV.check_pattern(many)[1]) and EV.check_pattern(many)[0] is None
    edge = base.copy(); edge[0] = np.nan
    assert EV.check_pattern(edge)[0] is None
    inf = base.copy(); inf[50] = np.inf
    assert EV.check_pattern(inf)[1]["filled_deg"] == [float(engine.ANGLES[50])]
    assert EV.check_pattern(base) [1] is None

# ------------------------------------------------------------------ F03 terrain content and engine in the run key
@test
def f03_changed_terrain_is_recalculated():
    slug, proj = setup_study("Cache")
    r1, res1 = evaluate(slug)
    assert r1["runs"] > 0 and r1["reused"] == 0 and r1["failed"] == 0
    assert all(spec["engine"] == engine.engine_info()["id"] and len(spec["terrain_sha256"]) == 16 for spec in res1["runs"].values())
    r2, _ = evaluate(slug); assert r2["reused"] == r2["runs"], r2           # nothing changed: everything reused
    f = P.site_dir(slug, "s1") / "AZI-30.00.PRO"                              # same name, different ground
    lines = f.read_text().splitlines(); hdr, rows = lines[0], lines[1:]
    f.write_text("\n".join([hdr] + [f"{r.split()[0]} {float(r.split()[1]) + 15 * i}" for i, r in enumerate(rows)]) + "\n")
    r3, res3 = evaluate(slug)
    assert r3["reused"] < r3["runs"], r3
    old = {c["band"] + c["target"]: c["hfta_fom"] for c in res1["cells"]}; new = {c["band"] + c["target"]: c["hfta_fom"] for c in res3["cells"]}
    assert old != new, "FOM did not change after the terrain changed"
    assert res3["engine"]["name"] == "OpenHFTA"

# ------------------------------------------------------------------ F14 a tower that serves no target
@test
def f14_no_targets_means_none():
    slug, proj = setup_study("Serves")
    proj["towers"].append({**json.loads(json.dumps(proj["towers"][0])), "id": "w2", "name": "Unused", "targets": []})
    put(slug, proj)
    _, res = evaluate(slug)
    assert not [c for c in res["cells"] if c["tower"] == "w2"]
    assert [c for c in res["cells"] if c["tower"] == "w1"]                    # no 'targets' key = every target

# ------------------------------------------------------------------ F15 missing arrival-angle data
@test
def f15_missing_prn_data_is_reported():
    slug, proj = setup_study("NoData", bands=("20", "160"))                   # the synthetic file has no 160 m column
    proj["targets"].append({"id": "t2", "name": "NoFile", "prn": "", "heading": 30, "offsets": [0], "weights": [1], "bands": ["20"]})
    put(slug, proj)
    _, res = evaluate(slug)
    by = {(c["target"], c["band"]): c for c in res["cells"]}
    assert by[("t1", "20")]["hfta_fom"] is not None
    assert by[("t1", "160")]["hfta_fom"] is None and "no usable data" in by[("t1", "160")]["issues"][0]
    assert by[("t2", "20")]["hfta_fom"] is None and "No elevation-angle" in by[("t2", "20")]["issues"][0]
    for i, c in enumerate(res["cells"]):
        assert C.get(f"/api/v2/p/{slug}/cell/{i}").status_code == 200
    assert C.get(f"/api/v2/p/{slug}/cell/999").status_code == 404
    for kind in ("xlsx", "csv", "html", "replay"):
        assert C.get(f"/api/v2/p/{slug}/export/{kind}").status_code == 200, kind

# ------------------------------------------------------------------ F11 flat ground replaces the whole set
@test
def f11_flat_replaces_every_profile():
    slug, proj = new_project("Flat"); proj["sites"] = [{"id": "s1", "name": "S"}]; put(slug, proj)
    d = P.site_dir(slug, "s1"); d.mkdir(parents=True, exist_ok=True)
    (d / "ZZZ-0.PRO").write_text("meters\n0 100\n30 90\n60 80\n")
    assert C.post(f"/api/v2/p/{slug}/sites/s1/flat").status_code == 200
    idx = TR.profile_index(d)
    assert len(idx) == 360 and all(v.startswith("FLAT") for v in idx.values())

@test
def f11_import_replaces_same_heading():
    d = TMP / "imp"; d.mkdir()
    (d / "OLD-AZI-30.00.PRO").write_text("meters\n0 0\n30 1\n")
    r = TR.import_files(d, [("NEW-30.00.PRO", b"meters\n0 0\n30 2\n")])
    assert r["replaced"] == ["OLD-AZI-30.00.PRO"] and TR.profile_index(d) == {30.0: "NEW-30.00.PRO"}

# ------------------------------------------------------------------ F12 old terrain kept when generation fails
@test
def f12_failed_generation_keeps_old_profiles():
    d = TMP / "gen"; d.mkdir(); (d / "KEEP-AZI-0.00.PRO").write_text("meters\n0 0\n30 1\n")
    class FailingDEM:
        n = 0
        def sample(self, la, lo):
            self.n += 1
            if self.n > 3: raise IOError("disk read error")
            return np.zeros(len(np.atleast_1d(la)))
    try: TR.generate_from_dem(FailingDEM(), 43.6, -72.0, d, "NEW", az_step=10)
    except IOError: pass
    else: raise AssertionError("no error")
    assert [p.name for p in d.glob("*.PRO")] == ["KEEP-AZI-0.00.PRO"]
    assert not [p for p in d.parent.glob(".stage-*")]
    for bad in (dict(az_step=0), dict(spacing_m=-1)):
        try: TR.generate_from_dem(FailingDEM(), 43.6, -72.0, d, "NEW", **bad)
        except ValueError: continue
        raise AssertionError(bad)

# ------------------------------------------------------------------ F16 GeoTIFF pixel-is-point
@test
def f16_pixel_is_point():
    import tifffile
    z = (np.arange(5)[:, None] * 10 + np.arange(5)[None, :]).astype(np.float32)      # z[row, col] = 10*row + col
    f = TMP / "pt.tif"
    geokeys = (1, 1, 0, 3, 1024, 0, 1, 2, 2048, 0, 1, 4326, 1025, 0, 1, 2)          # geographic, RasterPixelIsPoint
    tifffile.imwrite(f, z, extratags=[(33550, 12, 3, (1.0, 1.0, 0.0)), (33922, 12, 6, (0, 0, 0, 100.0, 200.0, 0)), (34735, 3, len(geokeys), geokeys)])
    dem = TR.DEM(f)
    assert abs(float(dem.sample([199.0], [101.0])[0]) - 11.0) < 1e-6, dem.sample([199.0], [101.0])

# ------------------------------------------------------------------ F07 + F08 Stack Designer search
def fake_search(groups_body, score_fn, min_gap, order="any", guys=None, radii=None):
    """Run CM.search with a fake scorer: score_fn(group index, heights) gives the FOM. No core needed."""
    proj = {"settings": {"spacing": dict(P.SPACING_DEFAULTS), "band_freq": {}}, "sites": [{"id": "s1"}],
            "targets": [{"id": "t1", "name": "A", "prn": "x.PRN", "heading": 0, "bands": ["20", "40"], "offsets": [0], "weights": [1]}],
            "towers": [{"id": "g", "guys": guys}] if guys else []}
    orig = (CM.Scorer.__init__, CM.Scorer.needed, CM.Scorer.compute, CM.Scorer.score, CM.aoa_status)
    def init(self, *a): self.idx = {0: "x"}; self.pat = {}; self.n_fom = 0; self.stages = []; self.n_failed = 0; self.n_filled = 0
    def score(self, model, bands, c, mode): return {"t1": {b: (score_fn(bands, c), list(c)) for b in bands}}
    CM.Scorer.__init__, CM.Scorer.needed, CM.Scorer.compute, CM.Scorer.score = init, lambda *a: [], lambda *a: None, score
    CM.aoa_status = lambda *a: ({1: 1.0}, None)
    try:
        body = {"site": "s1", "targets": ["t1"], "groups": groups_body, "min_gap": min_gap, "max_height": 300, "min_spacing": 1, "order": order, "top": 5}
        if guys: body["guys_from_tower"] = "g"
        return CM.search(proj, "x", body, None, {"cancel": False})
    finally:
        CM.Scorer.__init__, CM.Scorer.needed, CM.Scorer.compute, CM.Scorer.score, CM.aoa_status = orig

@test
def f07_constraint_aware_search_finds_feasible_layout():
    # Group A 1..100 ft, group B fixed at 100 ft, minimum gap 61 ft: only A <= 39 is allowed. Scores rise with height.
    r = fake_search([{"name": "A", "el": {"20": 4}, "count": 1, "lo": 1, "hi": 100, "step": 1},
                     {"name": "B", "el": {"40": 4}, "count": 1, "lo": 100, "hi": 100, "step": 1}],
                    lambda bands, c: sum(c) / 10, min_gap=61)
    assert r["rows"][0]["layout"] == [[39.0], [100.0]], r["rows"][0]["layout"]
    assert r["stats"]["complete"] is True

@test
def f07_matches_exhaustive_search():
    import itertools, random
    rnd = random.Random(7); table = {}
    sc = lambda bands, c: table.setdefault((bands[0], c), rnd.uniform(5, 15))
    groups = [{"name": "Tri", "el": {"20": 4}, "count": 2, "lo": 30, "hi": 150, "step": 10},
              {"name": "40", "el": {"40": 3}, "count": 1, "lo": 30, "hi": 150, "step": 10}]
    r = fake_search(groups, sc, min_gap=15, order="first_top")
    # brute force with the same rules
    best = None
    tri = [c for c in itertools.combinations([float(h) for h in range(30, 151, 10)], 2) if c[1] - c[0] >= 1]
    for a in tri:
        for b in [(float(h),) for h in range(30, 151, 10)]:
            if any(abs(x - y) < 15 for x in a for y in b) or min(a) <= max(b): continue
            v = (table[("20", a)] + table[("40", b)]) / 2
            if best is None or v > best[0]: best = (v, [list(a), list(b)])
    assert abs(r["rows"][0]["overall"] - round(best[0], 3)) < 1e-9 and r["rows"][0]["layout"] == best[1], (r["rows"][0], best)
    # a target's own best can never be worse than any shown layout for that target
    for row in r["rows"]: assert row["targets"]["t1"] <= r["own_best"]["t1"] + 1e-9

@test
def f08_guy_check_by_geometry_not_name():
    # Guy at 40 ft to 80 ft anchors. Radius 5 ft fits below it (up to about 35 ft); radius 60 ft fits only above 42 ft.
    # Two groups with the SAME name: the second must not reuse the first group's decisions. Scores favour low heights.
    guys = {"levels": [{"h": 40, "r": 80, "wire": "Phillystran HPTG-11200I", "tension": 1120}], "ice_in": 0, "margin_ft": 2, "density": 56}
    r = fake_search([{"name": "same", "el": {"20": 4}, "count": 1, "lo": 10, "hi": 98, "step": 2, "turning_radius": 5, "droop": 0},
                     {"name": "same", "el": {"40": 4}, "count": 1, "lo": 10, "hi": 98, "step": 2, "turning_radius": 60, "droop": 0}],
                    lambda bands, c: -sum(c) / 10, min_gap=1, guys=guys)
    assert r["rows"][0]["layout"][1][0] >= 42, r["rows"][0]["layout"]
    for row in r["rows"]:
        for (h,), radius in zip(row["layout"], (5, 60)):
            cl = G.check([{"name": "x", "h": h, "turning_radius": radius, "droop": 0}], guys["levels"], 0, 2, 56)[0]["clearance"]
            assert cl >= 0, (row["layout"], radius, cl)

# ------------------------------------------------------------------ F17 spreadsheet text stays text
@test
def f17_names_are_text_in_exports():
    slug, proj = setup_study("Formula")
    proj["sites"][0]["name"] = "=1+1"; put(slug, proj); evaluate(slug)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(C.get(f"/api/v2/p/{slug}/export/xlsx").data))
    cell = wb["Sites"]["A2"]; assert cell.value == "=1+1" and cell.data_type == "s", (cell.value, cell.data_type)
    assert any(c.data_type == "f" for row in wb["Results"].iter_rows() for c in row)          # own formulas still formulas
    csv_text = C.get(f"/api/v2/p/{slug}/export/csv").data.decode()
    assert "'=1+1" in csv_text

# ------------------------------------------------------------------ F18 cross-site requests refused
@test
def f18_cross_site_post_refused():
    slug, proj = new_project("Csrf"); proj["sites"] = [{"id": "s1", "name": "S"}]; put(slug, proj)
    assert C.post(f"/api/v2/p/{slug}/sites/s1/flat", headers={"Origin": "https://evil.example", "Sec-Fetch-Site": "cross-site"}).status_code == 403
    assert C.post(f"/api/v2/p/{slug}/sites/s1/flat", headers={"Origin": "https://evil.example"}).status_code == 403
    assert not TR.profile_index(P.site_dir(slug, "s1"))
    assert C.post(f"/api/v2/p/{slug}/sites/s1/flat", headers={"Origin": "http://localhost", "Sec-Fetch-Site": "same-origin"}).status_code == 200

# ------------------------------------------------------------------ L03 shareable exports leave the .PRN files out
@test
def l03_shareable_exports_have_no_prn():
    slug, proj = setup_study("Share"); evaluate(slug)
    z = zipfile.ZipFile(io.BytesIO(C.get(f"/api/v2/p/{slug}/export/replay").data))
    assert not [n for n in z.namelist() if n.lower().endswith(".prn") and n.startswith("prn/")]
    man = json.loads(z.read("manifest.json")); assert man["pack"] == "shareable" and man["prn_files"][0]["sha256"] == EX.sha256_file(PRN)
    assert man["engine"]["name"] == "OpenHFTA"
    z = zipfile.ZipFile(io.BytesIO(C.get(f"/api/v2/p/{slug}/export/replay?prn=1").data)); assert "prn/SYN.PRN" in z.namelist()
    z = zipfile.ZipFile(io.BytesIO(C.get(f"/api/v2/p/{slug}/export/bundle?prn=0").data))
    assert not [n for n in z.namelist() if n.startswith("prn/")] and "prn-files.txt" in z.namelist()
    z = zipfile.ZipFile(io.BytesIO(C.get(f"/api/v2/p/{slug}/export/bundle").data)); assert "prn/SYN.PRN" in z.namelist()
    # the private bundle opens again as a new project
    r = C.post("/api/v2/projects/import", data={"file": (io.BytesIO(C.get(f"/api/v2/p/{slug}/export/bundle").data), "b.zip")}, content_type="multipart/form-data")
    assert r.status_code == 200 and P.prn_file(r.get_json()["slug"], "SYN.PRN").exists()

# ------------------------------------------------------------------ F21 atomic writes
@test
def f21_atomic_write_leaves_no_partial_file():
    f = TMP / "a.json"; P.write_atomic(f, '{"a": 1}'); P.write_atomic(f, '{"a": 2}')
    assert json.loads(f.read_text()) == {"a": 2} and not list(TMP.glob(".a.json*"))

# ------------------------------------------------------------------ F21 revisions: two windows
@test
def f21_second_window_cannot_overwrite():
    slug, proj = new_project("Rev")
    a = json.loads(json.dumps(proj)); b = json.loads(json.dumps(proj))        # two browser windows with the same copy
    a["name"] = "from window A"; r = put(slug, a); assert r.status_code == 200 and r.get_json()["project"]["rev"] == a["rev"] + 1
    b["name"] = "from window B"; r = put(slug, b); assert r.status_code == 409, r.status_code
    assert "another window" in r.get_json()["error"] and P.load(slug)["name"] == "from window A"
    a2 = r.get_json()["project"]; a2["name"] = "B after reload"; assert put(slug, a2).status_code == 200
    legacy = json.loads(json.dumps(P.load(slug))); legacy.pop("rev"); legacy["name"] = "old client"
    assert put(slug, legacy).status_code == 200                                  # a client without rev still saves

@test
def csp_header_present():
    r = C.get("/"); assert "script-src 'self'" in r.headers.get("Content-Security-Policy", "") and r.headers.get("X-Frame-Options") == "DENY"
    assert "<script>" not in C.get("/classic").data.decode()                     # no inline script on the classic page

# ------------------------------------------------------------------ guy model against an exact catenary
def _catenary_sag(H, R, r, wl, d, tension, ice_in, density=56.0):
    """Exact inclined catenary from (0, H) to (R, 0), weight per foot wl + ice, cable tension at mid-span = tension.
    Independent of core/guys.py (closed-form shape, bisection on the horizontal tension)."""
    wi = math.pi * ((d / 2 + ice_in) ** 2 - (d / 2) ** 2) / 144 * density if ice_in > 0 else 0.0
    w = wl + wi
    def shape(Th):
        a = Th / w; s = -H / (2 * a * math.sinh(R / (2 * a))); u = (R - 2 * a * math.asinh(s)) / 2
        c = -a * math.cosh((R - u) / a); return (lambda x: a * math.cosh((x - u) / a) + c), u, a
    def mid_tension(Th):
        y, u, a = shape(Th); return Th * math.cosh((R / 2 - u) / a)
    lo, hi = w * R / 500, tension
    for _ in range(100):
        m = (lo + hi) / 2
        if mid_tension(m) > tension: hi = m
        else: lo = m
    y, u, a = shape((lo + hi) / 2)
    return H * (1 - r / R) - y(r)

@test
def guys_parabola_matches_catenary():
    worst = 0.0
    for H, R, ice in ((50, 100, 1.0), (100, 80, 1.25), (150, 120, 0.5), (40, 150, 1.0), (120, 60, 1.25), (100, 100, 0.0)):
        for wire, (wl, d, T) in G.WIRES.items():
            for r in (10, 25, 40):
                if r >= R: continue
                worst = max(worst, abs(_catenary_sag(H, R, r, wl, d, T, ice) - G.sag_at(H, R, r, wire, T, ice)))
    assert worst < 0.75, worst
    # geometry: zero anchor radius refused; a turning radius past the anchor is reported, not computed
    try: G.check([{"name": "a", "h": 50, "turning_radius": 20, "droop": 0}], [{"h": 60, "r": 0, "wire": "1/4 EHS steel", "tension": 670}])
    except ValueError: pass
    else: raise AssertionError("zero anchor radius accepted")
    row = G.check([{"name": "a", "h": 50, "turning_radius": 90, "droop": 0}], [{"h": 60, "r": 80, "wire": "1/4 EHS steel", "tension": 670}])[0]
    assert row["clearance"] is None and "anchor" in row["limit"]
    # no ice: the antenna below a 60 ft guy with 80 ft anchors, 20 ft out: straight wire at 45 ft, minus the wire's own
    # small sag; antenna at 40 ft, margin 2 -> a little under 3 ft
    row = G.check([{"name": "a", "h": 40, "turning_radius": 20, "droop": 0}], [{"h": 60, "r": 80, "wire": "1/4 EHS steel", "tension": 670}], ice_in=0, margin=2)[0]
    assert 2.6 < row["clearance"] <= 3.0, row

# ------------------------------------------------------------------ worldwide terrain from .hgt tiles
@test
def hgt_tiles_sample_and_generate():
    """Synthetic 1201-node tiles (3 arc-second): z = 100 * lat + 10 * lon (in degrees) so every sample is checkable,
    and a tile edge inside the profile set."""
    n = 1201
    def make(la, lo):
        lat = la + 1 - np.arange(n) / (n - 1); lon = lo + np.arange(n) / (n - 1)
        z = (100 * lat[:, None] + 10 * lon[None, :]).round().astype(">i2"); return z.tobytes()
    served = {}
    def loader(la, lo):
        served[(la, lo)] = served.get((la, lo), 0) + 1
        return make(la, lo) if (la, lo) in ((51, 6), (51, 7), (50, 6), (50, 7)) else None
    dem = TR.TerrainTiles(loader)
    z = dem.sample([51.5, 51.0, 50.9999, 51.999], [6.5, 7.0, 6.5, 7.9])
    assert np.allclose(z, [5215.0, 5170.0, 5165.0, 5278.9], atol=0.6), z
    assert np.isnan(dem.sample([49.5], [6.5])[0])                   # no tile there
    assert TR.tile_name(51, 6) == "N51E006" and TR.tile_name(-34, -59) == "S34W059"
    assert TR.tile_url(55, -132).endswith("/skadi/N55/N55W132.hgt.gz")
    assert set(TR.tiles_needed(51.02, 6.98, 4410)) == {(50, 6), (50, 7), (51, 6), (51, 7)}
    d = TMP / "tiles"; d.mkdir()
    r = TR.generate_from_dem(dem, 51.02, 6.98, d, "DL", az_step=45)   # profiles cross both tile edges
    assert r["written"] == 8 and not r["short_azimuths"], r
    idx = TR.profile_index(d); dist, h = TR.load_profile(d, idx[0.0])
    assert abs(h[0] / 3.28084 - (100 * 51.02 + 10 * 6.98)) < 1 and h[-1] > h[0]   # north: higher lat, higher z
    assert all(v == 1 for v in served.values())                      # each tile loaded once
    assert TR.in_usa(43.67, -72.01) and TR.in_usa(55.3, -131.6) and not TR.in_usa(51.2, 6.8)

@test
def hgt_fetch_route_uses_cache():
    n = 1201; z = (np.full((n, n), 250)).astype(">i2").tobytes()
    calls = []
    orig = TR.download_tile
    def fake(la, lo, cache_dir, timeout=180):
        calls.append((la, lo)); cache_dir.mkdir(parents=True, exist_ok=True)
        f = cache_dir / f"{TR.tile_name(la, lo)}.hgt"
        if not f.exists(): f.write_bytes(z)
        return f.read_bytes()
    TR.download_tile = fake
    try:
        slug, proj = new_project("Tiles"); proj["sites"] = [{"id": "s1", "name": "DL", "lat": 51.23, "lon": 6.78}]; put(slug, proj)
        r = C.post(f"/api/v2/p/{slug}/sites/s1/auto", json={"lat": 51.23, "lon": 6.78, "prefix": "DL", "az_step": 30}).get_json()
        assert r["written"] == 12 and r["source"] == TR.TERRAIN_TILES_SOURCE and r["ground_elevation_m"] == 250 and r["resolution_m"] > 90, r
        assert (P.ROOT / "_cache" / "hgt" / "N51E006.hgt").exists()
        r2 = C.post(f"/api/v2/p/{slug}/sites/s1/auto", json={"lat": 51.23, "lon": 6.78, "prefix": "DL", "az_step": 30, "source": "tiles"}).get_json()
        assert r2["written"] == 12
    finally:
        TR.download_tile = orig

@test
def hgt_bad_tile_rejected():
    try: TR.HGT(b"\x00" * 1000, 0, 0)
    except ValueError: pass
    else: raise AssertionError("accepted a tile of the wrong size")

# ------------------------------------------------------------------ VOACAP angle tables (core/aoa_gen.py)
@test
def voacap_deck_columns():
    from core import aoa_gen as AG
    d = AG.deck((42.4, -71.1), (-33.9, 151.2), [1, 11], 70, [3.6, 14.2], "W1 to Sydney")
    lines = d.splitlines()
    cir = next(l for l in lines if l.startswith("CIRCUIT"))
    assert cir == "CIRCUIT   42.40N    71.10W    33.90S   151.20E  S     0", cir
    assert next(l for l in lines if l.startswith("SUNSPOT")) == "SUNSPOT     70."
    assert [l for l in lines if l.startswith("MONTH")] == ["MONTH      2026 1.00", "MONTH      202611.00"]
    assert next(l for l in lines if l.startswith("FREQUENCY")) == "FREQUENCY  3.6014.20 0.00 0.00 0.00 0.00 0.00 0.00 0.00 0.00 0.00"
    assert d.count("EXECUTE") == 2 and lines[-1] == "QUIT"

@test
def voacap_parse_all_modes():
    from core import aoa_gen as AG
    rows = list(AG.parse_method25((DATA / "voacap_method25_sample.out").read_text()))
    hours = {r[0] for r in rows}; freqs = {r[2] for r in rows}
    assert len(hours) == 24 and {3.6, 14.2} <= freqs, (hours, freqs)
    h1 = next(r for r in rows if r[0] == 1.0 and r[2] == 3.6)
    assert h1[1] == "Jan" and len(h1[3]) == 5 and abs(h1[3][0]["angle"] - 6.58) < 0.01 and h1[3][0]["rel"] == 1.0
    # to_prn: percentages per band add up to 100 where there is data; angles are whole degrees from 1
    counts = {"80": {7: 2.0, 14: 1.0, 21: 1.0}, "20": {}}
    txt = AG.to_prn({"counts": counts, "weight_total": {"80": 4.0, "20": 0.0}}, "test", bands=["80", "20"], max_angle=25)
    lines = txt.splitlines(); assert lines[0].startswith("VOACAP-generated: test") and lines[1].split() == ["Elev", "80m", "20m"]
    vals = {int(l.split()[0]): [float(x) for x in l.split()[1:]] for l in lines[2:]}
    assert vals[7] == [50.0, 0.0] and vals[14] == [25.0, 0.0] and abs(sum(v[0] for v in vals.values()) - 100) < 0.01
    assert engine.parse_aoa(txt, 3.6)[7] == 0.5          # OpenHFTA reads the file

@test
def voacap_generate_if_installed():
    from core import aoa_gen as AG
    if not AG.available(): print("     (voacapl not installed: generation not tested)"); return
    slug, proj = new_project("Voacap")
    proj["sites"] = [{"id": "s1", "name": "Home", "lat": 42.4, "lon": -71.1}]; put(slug, proj)
    body = {"lat": 42.4, "lon": -71.1, "region": "JA", "name": "Home-JA-VOACAP.PRN", "months": [1, 7], "ssn": [70], "origin_name": "Home"}
    jid = C.post(f"/api/v2/p/{slug}/prn/generate", json=body).get_json()["job"]
    for _ in range(600):
        j = C.get(f"/api/job/{jid}").get_json()
        if j["state"] != "running": break
        time.sleep(0.3)
    assert j["state"] == "done", j
    f = P.prn_file(slug, "Home-JA-VOACAP.PRN"); assert f.exists() and "20" in j["result"]["bands"]
    txt = f.read_text(); assert txt.startswith("VOACAP-generated: Home to Japan (months Jan,Jul, all hours, SSN 70)"), txt[:90]
    aoa = engine.parse_aoa(txt, 14.2); assert aoa and abs(sum(aoa.values()) - 1) < 0.01
    assert sum(a * p for a, p in aoa.items()) < 12          # Japan from New England arrives low
    jid = C.post(f"/api/v2/p/{slug}/prn/generate", json={**body, "name": "../x.PRN"}).get_json()["job"]
    time.sleep(1); assert C.get(f"/api/job/{jid}").get_json()["state"] == "error"

if __name__ == "__main__":
    print(f"\nPASS: {len(PASSED)} tests")
    if APP.POOL is not None: APP.POOL.shutdown(wait=True, cancel_futures=True)
